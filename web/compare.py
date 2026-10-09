"""ブラウザ版の結果が、デスクトップ版と同じ Python の処理の結果と一致するかを確かめる。

  node web/test_engine.mjs dist-web 出力フォルダ    # 先に、ブラウザ版の中身（Pyodide）で作る
  python web/compare.py dist-web 出力フォルダ        # 同じ入力をこちらの Python で処理して比べる

読み取り結果（番号・状態・警告・向き・用紙の種類）がすべて同じで、PDFを画像にしたときに画素が一致し、
香盤表（Excel）のすべてのセルが同じなら合格。
"""
import json, os, sys, tempfile
import numpy as np
import openpyxl
import pypdfium2 as pdfium

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from ekonte.omr import read_pdf
from ekonte.compose import compose, make_cut
from ekonte.sheet import build as build_sheet

FIELDS = ("page", "name", "status", "warnings", "orientation", "format")


def render(path):
    pdf = pdfium.PdfDocument(path)
    return [np.asarray(pdf[i].render(scale=2).to_pil().convert("L"), dtype=np.int16) for i in range(len(pdf))]


def pdf_diff(a, b):
    """2つのPDFの違い（同じなら None）"""
    x, y = render(a), render(b)
    if len(x) != len(y):
        return f"ページ数が違う（{len(x)} / {len(y)}）"
    for i, (p, q) in enumerate(zip(x, y)):
        if p.shape != q.shape:
            return f"p{i + 1} の大きさが違う"
        worst = int(np.abs(p - q).max())
        if worst:
            return f"p{i + 1} の画素が違う（最大差 {worst}、{int((p != q).sum())}画素）"
    return None


def xlsx_diff(a, b):
    """2つの Excel のシートとセルの違い（同じなら None）"""
    x, y = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    if x.sheetnames != y.sheetnames:
        return f"シートが違う（{x.sheetnames} / {y.sheetnames}）"
    for name in x.sheetnames:
        p, q = list(x[name].iter_rows(values_only=True)), list(y[name].iter_rows(values_only=True))
        if p != q:
            return f"「{name}」のシートの中身が違う"
    return None


def main(dist, out):
    with open(os.path.join(out, "results.json"), encoding="utf-8") as f:
        rec = json.load(f)
    samples = os.path.join(dist, "samples")
    book = os.path.join(samples, "香盤表・カット表_記入例.xlsx")
    fails = []
    if not rec["info"]["excel_guard"]:
        fails.append("ブラウザ版で、悪意のあるExcelへの対策（defusedxml）が効いていない")
    for label, msg in rec["errors"].items():
        if not msg:
            fails.append(f"{label}: エラーにならなかった")
    if not rec["keptOk"]:
        fails.append("Excelを読み込めなかったあと、前に選んだExcelで書き出せなかった")
    with tempfile.TemporaryDirectory() as tmp:
        for case in rec["cases"]:
            native = list(read_pdf(os.path.join(samples, case["pdf"])))
            got = [{k: getattr(r, k) for k in FIELDS} for r in native]
            want = [{k: p[k] for k in FIELDS} for p in case["pages"]]
            if got != want:
                fails.append(f"{case['pdf']}: 読み取り結果が違う\n  ブラウザ版 {want}\n  Python    {got}")
            if not case["previews"]:
                fails.append(f"{case['pdf']}: 画面に出す画像が届いていないページがある")
            cuts = [make_cut(name, native[i].picture) for i, name in case["rows"]]
            mine = os.path.join(tmp, case["output"])
            mine_kouban = os.path.join(tmp, case["kouban"])
            compose(cuts, mine, book, case["format"], kouban_out=mine_kouban)
            for name, diff in ((case["output"], pdf_diff(os.path.join(out, case["output"]), mine)),
                               (case["kouban"], xlsx_diff(os.path.join(out, case["kouban"]), mine_kouban))):
                print(f"{name}: {diff or '一致'}")
                if diff:
                    fails.append(f"{name}: {diff}")
        for s in rec["sheets"]:
            mine = os.path.join(tmp, s["output"])
            build_sheet(mine, s["pages"], s["fmt"])
            diff = pdf_diff(os.path.join(out, s["output"]), mine)
            print(f"{s['output']}: {diff or '一致'}")
            if diff:
                fails.append(f"{s['output']}: {diff}")
    if fails:
        print("\n".join(["✗ ブラウザ版とデスクトップ版で結果が違います"] + fails))
        sys.exit(1)
    print("✓ ブラウザ版とデスクトップ版の結果は一致しました")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
