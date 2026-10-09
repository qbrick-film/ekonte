"""絵コンテと一緒に書き出す香盤表（Excel）。絵コンテと同じ中身（カット絵以外すべて）を2枚のシートにする。

  香盤表     1行1シーン。元の香盤表（場面・L/S・D/N・ロケ地・登場人物など・備考）に、
             カット表から集めたカット数・時間・小道具・カメラ・機材を足す。最後の行に合計
  カット一覧 カット表と同じ形（ACTION・SE・SCENARIO・TIME・小道具・カメラ・機材・NOTE の列、書いた行のまま）。
             シーンの頭にシーンの行（場面・L/S・D/N・ロケ地・備考と、シーン全体の小道具など）。カット絵だけのカットも載せる

中身は compose.plan() の結果から作るので、絵コンテPDFと食い違わない。
直すときは元の Excel（香盤表・カット表）を直して書き出し直す（このファイルは読み込めないようにしてある）。
"""
import os, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from .compose import without
from .excel_import import OUTPUT_MARK

THIN = Side(style="thin", color="999999")
BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
SCENE_TOP = Border(top=Side(style="medium", color="555555"), bottom=THIN, left=THIN, right=THIN)  # シーンの区切り
HEAD_FILL = PatternFill("solid", fgColor="DDDDDD")
TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")
SCENE_FILL = PatternFill("solid", fgColor="EAF1FB")  # シーンの行（カット表のテンプレートと同じ色）
CONT_FILL = PatternFill("solid", fgColor="F5F5F5")   # 続きの行
LEFT = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")
GATHER = [("props", "小道具", 24), ("camera", "カメラ", 18), ("gear", "機材", 18)]  # カット表からシーンごとにまとめるもの


def output_path(pdf_path, source=None):
    """絵コンテPDFの隣に置く香盤表の名前（絵コンテ.pdf → 絵コンテ_香盤表.xlsx）。元の Excel と同じ名前にはしない。"""
    base = os.path.splitext(pdf_path)[0] + "_香盤表"
    if source and os.path.abspath(base + ".xlsx") == os.path.abspath(source):
        base += "_書き出し"
    return base + ".xlsx"


def write(path, rows, book, source=None):
    """rows: compose.plan() の結果。book: 読み込んだ香盤表・カット表。source: 元の Excel の場所（ファイル名を書き添える）"""
    wb = Workbook()
    _scenes(wb.active, rows, book, source)
    _cuts(wb.create_sheet("カット一覧"), rows, book)
    wb.properties.subject = OUTPUT_MARK
    wb.properties.creator = "絵コンテ作成ソフト"
    wb.save(path)


def _number(v):
    """数字だけの文字は数として書く（Excel で合計などに使えるように）。"""
    if isinstance(v, str) and re.fullmatch(r"\d+(\.\d+)?", v):
        f = float(v)
        return int(f) if f.is_integer() else f
    return v


def _seconds(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _scene_order(s):
    m = re.match(r"(\d+)(.*)", s)
    return (int(m.group(1)), m.group(2)) if m else (float("inf"), s)


def _cell(ws, row, col, value, align=LEFT, border=BORDER, **font):
    c = ws.cell(row, col, value)
    c.alignment, c.border = align, border
    if font:
        c.font = Font(**font)
    return c


def _head(ws, row, columns):
    for i, (name, width) in enumerate(columns, 1):
        c = _cell(ws, row, i, name, CENTER, bold=True)
        c.fill = HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = width


def _print_setup(ws, title_rows):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = title_rows


def _scenes(ws, rows, book, source):
    ws.title = "香盤表"
    gathered = {}  # S → {"cuts": カット数, "time": 秒の合計, "items": {キー: [...]}}
    for r in rows:
        g = gathered.setdefault(r.cut["s"], {"cuts": 0, "time": None, "items": {k: [] for k, _, _ in GATHER}})
        g["cuts"] += 1
        t = _seconds(r.time)
        if t is not None:
            g["time"] = (g["time"] or 0) + t
        for k, _, _ in GATHER:
            g["items"][k] += r.items.get(k, [])
    extra = [x for x in book.extra if x[2]]  # 香盤表のそのほかの列（登場人物など）。どのシーンにも書いていない列は省く
    scenes = set(gathered) | set(book.kouban) | {s for _, _, values in extra for s in values}

    ws["A1"] = "香盤表"
    ws["A1"].font = Font(bold=True, size=16)
    for row, (label, value) in enumerate((("タイトル", book.title), ("撮影日", book.date)), 2):
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 3, value)
    note = "絵コンテ作成ソフトで作成" + (f"（元の Excel: {os.path.basename(source)}）" if source else "")
    ws.cell(4, 1, note + "。直すときは元の Excel を直して、もう一度書き出してください").font = Font(size=9, color="666666")

    head = 6  # 見出しの行（すぐ上の行に役者名など）
    columns = [("#S", 5), ("場面", 20), ("L/S", 6), ("D/N", 6), ("ロケ地", 18)]
    columns += [(name, 8) for name, _, _ in extra]
    columns += [("カット数", 7), ("時間（秒）", 8)] + [(label, width) for _, label, width in GATHER] + [("備考", 36)]
    _head(ws, head, columns)
    for j, (_, above, _) in enumerate(extra):
        if above:
            _cell(ws, head - 1, 6 + j, above, CENTER, border=Border(), size=9)
    centered = {1, 3, 4} | set(range(6, 6 + len(extra) + 2))

    r, total_cuts, total_time = head + 1, 0, None
    for s in sorted(scenes, key=_scene_order):
        info = book.kouban.get(s, {})
        g = gathered.get(s)
        if g:
            items = g["items"]
        else:  # カットのないシーンでも、シーン全体に書いた小道具などは載せる
            items = {k: book.sources.get(k, ({}, {}))[0].get(s, []) for k, _, _ in GATHER}
        items = {k: list(dict.fromkeys(v)) for k, v in items.items()}
        shown = {x for v in items.values() for x in v}
        values = [_number(s)] + [info.get(k, "") for k in ("bamen", "ls", "dn", "loca")]
        values += [v.get(s, "") for _, _, v in extra]
        values += [g["cuts"] if g else 0, g["time"] if g else None]
        values += ["、".join(items[k]) for k, _, _ in GATHER]
        values.append(without(info.get("biko", ""), shown))  # 小道具などの欄と同じものは備考に重ねない
        for i, v in enumerate(values, 1):
            _cell(ws, r, i, v, CENTER if i in centered else LEFT)
        total_cuts += g["cuts"] if g else 0
        if g and g["time"] is not None:
            total_time = (total_time or 0) + g["time"]
        r += 1
    cut_col = 6 + len(extra)
    for i in range(1, len(columns) + 1):
        _cell(ws, r, i, None, CENTER).fill = TOTAL_FILL
    ws.cell(r, 2, "合計").font = Font(bold=True)
    ws.cell(r, cut_col, total_cuts).font = Font(bold=True)
    ws.cell(r, cut_col + 1, total_time).font = Font(bold=True)
    ws.freeze_panes = ws.cell(head + 1, 2)
    _print_setup(ws, f"{head}:{head}")


def _cuts(ws, rows, book):
    """カット表と同じ形の一覧。カットは絵コンテと同じ順（番号順）、各カットの行は書いたとおりに並べる。"""
    cams = book.cams or ["A"]
    columns = [("S", "s", 5), ("C", "c", 6), ("ACTION", "action", 30), ("SE", "se", 14), ("SCENARIO", "scenario", 34),
               ("TIME", "time", 7), ("小道具", "props", 16)]
    columns += [(f"カメラ{label}", ("cams", label), 11) for label in cams]
    columns += [("機材", "gear", 14), ("NOTE", "note", 26)]
    _head(ws, 1, [(name, width) for name, _, width in columns])
    scene_lines, cut_lines = book.lines
    by_scene = {}
    for r in rows:
        by_scene.setdefault(r.cut["s"], []).append(r)
    r_out, total = 2, None

    def put(line, fill, top):
        nonlocal r_out
        for j, (_, key, _) in enumerate(columns, 1):
            v = line.get("cams", {}).get(key[1], "") if isinstance(key, tuple) else line.get(key, "")
            c = _cell(ws, r_out, j, _number(v) if key in ("s", "c", "time") else (v or None),
                      CENTER if key in ("s", "c", "time") else LEFT, SCENE_TOP if top else BORDER)
            if fill:
                c.fill = fill
        r_out += 1

    for s in sorted(set(by_scene) | set(scene_lines) | set(book.kouban), key=_scene_order):
        # シーンの行：場面・L/S・D/N・ロケ地（ACTION の列）、カット表のシーン全体の行、香盤表の小道具・機材、備考（NOTE の列）
        info = book.kouban.get(s, {})
        head = " / ".join(v for v in (info.get("bamen"), info.get("ls"), info.get("dn")) if v)
        title = "　".join(x for x in (f"【{head}】" if head else "", f"ロケ地: {info['loca']}" if info.get("loca") else "") if x)
        lines = [dict(line) for line in scene_lines.get(s, [])]
        for k, values in book.kouban_items.get(s, {}).items():
            written = {line.get(k) for line in lines}
            lines += [{k: v} for v in values if v not in written]
        if title:
            if lines and not lines[0].get("action"):
                lines[0]["action"] = title
            else:
                lines.insert(0, {"action": title})
        items = {x for r in by_scene.get(s, []) for v in r.items.values() for x in v}
        items |= {x for k, (by_s, _) in book.sources.items() for x in by_s.get(s, [])}
        biko = without(info.get("biko", ""), items)  # 小道具などと同じものは備考に重ねない
        if biko:
            free = next((line for line in lines if not line.get("note")), None)
            if free is None:
                lines.append(free := {})
            free["note"] = f"備考: {biko}"
        for i, line in enumerate(lines or [{}]):
            put({**line, "s": s} if i == 0 else line, SCENE_FILL, i == 0)
        # カット：書いた行のまま（カット表にないカットは S・C だけ）
        for r in by_scene.get(s, []):
            cut = r.cut
            for i, line in enumerate(cut_lines.get((cut["s"], cut["c"])) or [{}]):
                put({**line, "s": cut["s"], "c": cut["c"]} if i == 0 else line, CONT_FILL if i else None, False)
            t = _seconds(r.time)
            if t is not None:
                total = (total or 0) + t
    for j in range(1, len(columns) + 1):
        _cell(ws, r_out, j, None, CENTER).fill = TOTAL_FILL
    ws.cell(r_out, 3, "合計").font = Font(bold=True)
    ws.cell(r_out, 6, total).font = Font(bold=True)
    ws.freeze_panes = "C2"
    _print_setup(ws, "1:1")
