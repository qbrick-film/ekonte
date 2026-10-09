"""香盤表・カット表・小道具・機材のExcelを読み込む。

見出し行は列名で自動検出するので、列の並びや開始行が変わっても読める。
"""
import re, zipfile
from collections import defaultdict
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

SCENE_HEADERS = ("#S", "S#", "シーン", "S")
CUT_HEADERS = ("カット", "C", "C#")


def _open(path):
    """Excelを開く。開けなければ、理由が分かる日本語のエラーにする。"""
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException) as e:
        raise ValueError("Excel（.xlsx）として開けません。古い形式（.xls）などのファイルなら、"
                         "Excel で .xlsx 形式で保存し直してから選んでください") from e


def _norm(v):
    return re.sub(r"\s", "", str(v)) if v is not None else ""


def _num(v):
    """シーン/カット番号を文字列に正規化する（2, 2.0, '２' → '2'）。"""
    if v is None or _norm(v) == "":
        return None
    s = _norm(v).translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    if s.isdigit():
        return str(int(s))
    m = re.fullmatch(r"0*(\d+)([A-Ja-j])", s)  # 追加カット "03A" → "3a"
    if m:
        return str(int(m.group(1))) + m.group(2).lower()
    return s


def _find_header(ws, required):
    """required のいずれかを含む行を見出し行とし、{見出し: 列番号} を返す。"""
    for row in ws.iter_rows(min_row=1, max_row=30):
        names = {_norm(c.value): c.column for c in row if c.value is not None}
        if any(h in names for h in required):
            return row[0].row, names
    raise ValueError(f"見出し行が見つかりません（{'/'.join(required)} の列が必要）")


def _col(names, *candidates, prefix=False):
    for c in candidates:
        for name, idx in names.items():
            if name == c or (prefix and name.startswith(c)):
                return idx
    return None


def load_kouban(path):
    """香盤表 → {シーン: {場面, ls, dn, loca, biko}}"""
    ws = _open(path).active
    header_row, names = _find_header(ws, SCENE_HEADERS)
    cols = dict(
        scene=_col(names, *SCENE_HEADERS),
        bamen=_col(names, "場面"),
        ls=_col(names, "L/S", "L/LS"),
        dn=_col(names, "D/N"),
        loca=_col(names, "ロケ地"),
        biko=_col(names, "備考", prefix=True),
    )
    scenes = {}
    for row in ws.iter_rows(min_row=header_row + 1):
        get = lambda k: row[cols[k] - 1].value if cols[k] else None
        s = _num(get("scene"))
        if s is None:
            continue
        info = {k: (str(get(k)).strip() if get(k) is not None else "") for k in ("bamen", "ls", "dn", "loca", "biko")}
        if any(info.values()):
            scenes[s] = info
    return scenes


def _load_by_cut(path, fields, combine=None):
    """「シーン | カット | 項目...」形式のExcelを読む。

    fields: {キー: 見出し候補のタプル}
    combine: {新キー: (キー, ...)} 同じ行の値をスペースで連結して1項目にする
    戻り値: {キー: (シーン単位 {S: [..]}, カット単位 {(S, C): [..]})}
    カット欄が空の行はシーン全体に適用する。
    """
    ws = _open(path).active
    header_row, names = _find_header(ws, SCENE_HEADERS)
    c_scene = _col(names, *SCENE_HEADERS)
    c_cut = _col(names, *CUT_HEADERS)
    cols = {k: _col(names, *cands, prefix=True) for k, cands in fields.items()}
    if not any(cols.values()):
        raise ValueError(f"{path}: 項目の列が見つかりません（{'/'.join(c for cs in fields.values() for c in cs)}）")
    combine = combine or {}
    result = {k: (defaultdict(list), defaultdict(list)) for k in list(fields) + list(combine)}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        s = _num(row[c_scene - 1])
        if s is None:
            continue
        cut = _num(row[c_cut - 1]) if c_cut else None
        values = {k: str(row[c - 1]).strip() if c and row[c - 1] is not None else "" for k, c in cols.items()}
        for k, keys in combine.items():
            values[k] = " ".join(values[x] for x in keys if values[x])
        for k, v in values.items():
            if v:
                by_scene, by_cut = result[k]
                (by_cut[(s, cut)] if cut else by_scene[s]).append(v)
    return result


def load_props(path):
    """小道具Excel → {"props": (シーン単位, カット単位)}"""
    return _load_by_cut(path, {"props": ("小道具", "アイテム", "品名")})


def load_equipment(path):
    """カメラ・レンズ・機材Excel → {"camera": .., "gear": ..}

    1行 = カメラ1台の設定。カメラとレンズは同じ行どうしを組にする（例: "A cam 35mm"）。
    """
    r = _load_by_cut(path, {
        "cam": ("カメラ",),
        "lens": ("レンズ",),
        "gear": ("機材", "特機", "必要機材"),
    }, combine={"camera": ("cam", "lens")})
    return {"camera": r["camera"], "gear": r["gear"]}


CUT_FIELDS = {  # キー: 見出しの候補（前方一致）
    "action": ("ACTION/SE", "ACTION", "芝居"),
    "se": ("SE", "効果音"),  # 絵コンテでは ACTION/SE 欄に「SE：〜」として入れる
    "scenario": ("SCENARIO", "セリフ", "台詞"),
    "time": ("TIME", "秒"),
    "props": ("小道具",),
    "gear": ("機材", "特機"),
    "note": ("NOTE", "メモ"),
}


def load_cut_table(path):
    """カット表 → (カットの一覧, NOTE用のデータ)

    1行 = 1カット。書き方のルール:
      S と C を書いた行          … そのカット
      S だけ書いて C が空欄の行  … シーン全体（小道具・カメラ・機材・NOTE をそのシーンの全カットに入れる）
      S も C も空欄の行          … すぐ上の行の続き（各欄に1行ずつ追記する）
    同じ S・C を2回書いた場合も追記する。

    戻り値:
      cuts:    {(S, C): {"action": [...], "scenario": [...], "time": [...]}}  表に書いた順
      sources: compose の NOTE_ITEMS と同じ形 {キー: (シーン単位 {S: [..]}, カット単位 {(S, C): [..]})}
    """
    wb = _open(path)
    ws = wb["カット表"] if "カット表" in wb.sheetnames else wb.worksheets[0]
    header_row, names = _find_header(ws, SCENE_HEADERS)
    c_scene = _col(names, *SCENE_HEADERS)
    c_cut = _col(names, *CUT_HEADERS)
    if c_cut is None:
        raise ValueError("C（カット）の列が見つかりません")
    cols = {k: _col(names, *cands, prefix=True) for k, cands in CUT_FIELDS.items()}
    # カメラは「カメラA」「カメラB」… の列。値はレンズ（例: 35mm）
    cams = [(name[len("カメラ"):], idx) for name, idx in names.items() if name.startswith("カメラ")]

    cuts = {}
    sources = {k: (defaultdict(list), defaultdict(list)) for k in ("props", "gear", "note", "camera")}
    target = None  # 直前の行の対象: ("scene", S) / ("cut", (S, C))
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        cell = lambda idx: row[idx - 1] if idx and idx <= len(row) else None
        raw_s, raw_c = cell(c_scene), cell(c_cut)
        s, c = _num(raw_s), _num(raw_c)
        if s is not None and not s.isdigit():
            continue  # 「→ S」などの説明行
        values = {k: str(cell(idx)).strip() for k, idx in cols.items() if cell(idx) not in (None, "")}
        values = {k: v for k, v in values.items() if v}
        cam_values = [(label, str(cell(idx)).strip()) for label, idx in cams if cell(idx) not in (None, "")]
        if s is None and c is None:
            if not values and not cam_values:
                continue  # 空行
            if target is None:
                continue  # 先頭の続き行は対象がないので無視
        elif s is not None and c is None:
            target = ("scene", s)
        elif s is not None:
            target = ("cut", (s, c))
            cuts.setdefault((s, c), {"action": [], "scenario": [], "time": []})
        else:
            continue  # C だけ書いた行は不正なので無視
        kind, key = target
        for k, v in values.items():  # CUT_FIELDS の順なので、同じ行なら ACTION → SE の順に入る
            if k == "se":
                if kind == "cut":
                    cuts[key]["action"].append("SE：" + re.sub(r"^SE\s*[：:]\s*", "", v, flags=re.I))
            elif k in ("action", "scenario", "time"):
                if kind == "cut":
                    cuts[key][k].append(_fmt_time(v) if k == "time" else v)
            else:
                (sources[k][1][key] if kind == "cut" else sources[k][0][key]).append(v)
        for label, lens in cam_values:
            item = f"{label} {lens}".strip()
            (sources["camera"][1][key] if kind == "cut" else sources["camera"][0][key]).append(item)
    return cuts, sources


def _fmt_time(v):
    """3.0 → "3"、1.5 → "1.5"。数字でなければそのまま。"""
    try:
        f = float(v)
    except ValueError:
        return v
    return str(int(f)) if f == int(f) else f"{f:g}"
